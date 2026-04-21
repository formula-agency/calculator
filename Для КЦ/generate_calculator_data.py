#!/usr/bin/env python3
"""Builds browser-ready JS data from local Excel sources."""

from __future__ import annotations

import json
import re
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

import openpyxl


ROOT = Path(__file__).resolve().parent
OUT_FILE = ROOT / "mortgage_calculator_data.js"
RATES_FILE_NAME = "\u0421\u0442\u0430\u0432\u043a\u0438 NEW.xlsx"


def to_float(value: Any) -> float | None:
    if isinstance(value, (int, float)):
        return float(value)
    if isinstance(value, str):
        txt = value.strip().replace(" ", "").replace(",", ".")
        if not txt:
            return None
        try:
            return float(txt)
        except ValueError:
            return None
    return None


def to_str(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


def parse_min_pv(raw: str) -> float | None:
    match = re.search(r"(\d+(?:[.,]\d+)?)", raw)
    if not match:
        return None
    try:
        return float(match.group(1).replace(",", "."))
    except ValueError:
        return None


def parse_credit_limit(special: str) -> dict[str, float | None]:
    txt = to_str(special).lower().replace(",", ".")
    match = re.search(r"(\d+(?:\.\d+)?)\s*\u043c\u043b\u043d", txt)
    if not match:
        return {"loan_min": None, "loan_max": None}
    try:
        amount = float(match.group(1)) * 1_000_000.0
    except ValueError:
        return {"loan_min": None, "loan_max": None}

    if "\u0434\u043e" in txt:
        return {"loan_min": None, "loan_max": amount}
    if (
        "\u0431\u043e\u043b\u0435\u0435" in txt
        or "\u0441\u0432\u044b\u0448\u0435" in txt
        or "\u043e\u0442" in txt
    ):
        return {"loan_min": amount, "loan_max": None}
    return {"loan_min": None, "loan_max": None}


def pick_workbooks() -> tuple[Path, Path]:
    rates = ROOT / RATES_FILE_NAME
    if not rates.exists():
        raise RuntimeError(
            f"Rates file '{RATES_FILE_NAME}' not found. "
            "Calculator is configured to use only this file."
        )

    assortment = None
    for xlsx in sorted(ROOT.glob("*.xlsx")):
        if xlsx.resolve() == rates.resolve():
            continue

        wb = openpyxl.load_workbook(xlsx, read_only=True, data_only=True)
        ws = wb[wb.sheetnames[0]]
        header = list(next(ws.iter_rows(min_row=1, max_row=1, values_only=True)))
        wb.close()
        header_set = {str(h).strip() for h in header if h is not None}
        if {
            "\u041c\u0424\u0426",
            "\u041e\u0431\u044a\u0435\u043a\u0442",
            "\u0420\u0430\u0439\u043e\u043d",
            "\u041d\u043e\u043c\u0435\u0440",
        }.issubset(header_set):
            assortment = xlsx

    if not assortment:
        raise RuntimeError("Could not detect assortment Excel file in current folder.")
    return assortment, rates


def build_assortment(path: Path) -> list[dict[str, Any]]:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    header = list(next(ws.iter_rows(min_row=1, max_row=1, values_only=True)))
    idx = {name: i for i, name in enumerate(header)}

    required = ["GUID", "Район", "Объект", "Номер", "МФЦ"]
    missing = [col for col in required if col not in idx]
    if missing:
        wb.close()
        raise RuntimeError(f"Assortment file is missing required columns: {missing}")

    def get(row: tuple[Any, ...], col: str) -> Any:
        i = idx.get(col)
        if i is None or i >= len(row):
            return None
        return row[i]

    result: list[dict[str, Any]] = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        mfc = to_float(get(row, "\u041c\u0424\u0426"))
        if mfc is None or mfc <= 0:
            continue
        public_price = to_float(get(row, "\u041f\u0443\u0431\u043b\u0438\u0447\u043d\u0430\u044f \u0446\u0435\u043d\u0430"))
        if public_price is None or public_price <= 0:
            public_price = mfc

        apartment_raw = get(row, "Номер")
        apartment = to_str(apartment_raw)
        if apartment_raw is not None and isinstance(apartment_raw, float) and apartment_raw.is_integer():
            apartment = str(int(apartment_raw))

        area = to_float(get(row, "Площадь"))
        record = {
            "id": to_str(get(row, "GUID")) or f"row-{len(result) + 1}",
            "status": to_str(get(row, "Статус")),
            "district": to_str(get(row, "Район")),
            "object": to_str(get(row, "Объект")),
            "section": to_str(get(row, "Подъезд (секция)")),
            "layout": to_str(get(row, "Тип планировки")),
            "floor": to_float(get(row, "Этаж")),
            "apartment": apartment,
            "area": area,
            "mfc": round(mfc, 2),
            "public_price": round(public_price, 2),
        }
        result.append(record)

    wb.close()
    result.sort(
        key=lambda x: (
            x["district"],
            x["object"],
            x["apartment"],
            x["id"],
        )
    )
    return result


def build_rates(path: Path) -> list[dict[str, Any]]:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    header = list(next(ws.iter_rows(min_row=1, max_row=1, values_only=True)))
    idx = {name: i for i, name in enumerate(header)}

    commission_col = None
    for name in header:
        if isinstance(name, str) and name.strip().startswith("Комиссия"):
            commission_col = name
            break
    if not commission_col:
        wb.close()
        raise RuntimeError("Rates file is missing commission column.")

    required = ["БАНК", "Особые условия", "Программа", "Срок кредита", "Срок действия дисконта", "Ставка для клиента", "Базовая ставка", "Размер ПВ"]
    missing = [col for col in required if col not in idx]
    if missing:
        wb.close()
        raise RuntimeError(f"Rates file is missing required columns: {missing}")

    def get(row: tuple[Any, ...], col: str) -> Any:
        i = idx.get(col)
        if i is None or i >= len(row):
            return None
        return row[i]

    result: list[dict[str, Any]] = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        commission = to_float(get(row, commission_col))
        client_rate = to_float(get(row, "Ставка для клиента"))
        if commission is None or client_rate is None:
            continue

        special = to_str(get(row, "Особые условия"))
        pv_text = to_str(get(row, "Размер ПВ"))
        limits = parse_credit_limit(special)

        rec = {
            "bank": to_str(get(row, "БАНК")),
            "special": special,
            "program": to_str(get(row, "Программа")),
            "loan_term_rule": to_str(get(row, "Срок кредита")),
            "discount_term": to_str(get(row, "Срок действия дисконта")),
            "discount_delta": to_float(get(row, "Размер дисконта к процентной ставке (в процентных пунктах)")),
            "base_rate": to_float(get(row, "Базовая ставка")),
            "client_rate": client_rate,
            "commission": commission,
            "pv_text": pv_text,
            "min_pv": parse_min_pv(pv_text),
            "loan_min": limits["loan_min"],
            "loan_max": limits["loan_max"],
        }
        if not rec["bank"] or not rec["program"] or not rec["discount_term"]:
            continue
        result.append(rec)

    wb.close()
    return result


def main() -> None:
    assortment_file, rates_file = pick_workbooks()
    assortment = build_assortment(assortment_file)
    rates = build_rates(rates_file)

    payload = {
        "generated_at_utc": datetime.now(timezone.utc).isoformat(),
        "sources": {
            "assortment_file": assortment_file.name,
            "rates_file": rates_file.name,
        },
        "counts": {
            "apartments": len(assortment),
            "rates": len(rates),
        },
        "assortment": assortment,
        "rates": rates,
    }

    OUT_FILE.write_text(
        "window.MORTGAGE_CALC_DATA = " + json.dumps(payload, ensure_ascii=False, separators=(",", ":")) + ";\n",
        encoding="utf-8",
    )
    print(f"Written: {OUT_FILE.name}")
    print(f"Apartments: {len(assortment)}")
    print(f"Rates: {len(rates)}")


if __name__ == "__main__":
    main()
